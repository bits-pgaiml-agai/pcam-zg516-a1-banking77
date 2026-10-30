"""Build a single Word speaking script for the Group 15 demo (3 members)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo" / "Group15_Demo_Speaking_Script.docx"


def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def p(doc, text, *, size=11, bold=False, italic=False, after=6, align=None):
    para = doc.add_paragraph()
    if align:
        para.alignment = align
    run = para.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.space_before = Pt(0)
    return para


def speak(doc, text):
    """Spoken lines — slightly larger and clear."""
    para = doc.add_paragraph()
    run = para.add_run(text)
    font(run, size=12)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.left_indent = Inches(0.15)
    return para


def cue(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    font(run, size=10, italic=True, color=RGBColor(0x55, 0x55, 0x55))
    para.paragraph_format.space_after = Pt(6)
    return para


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        font(r, size=10, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            font(r, size=10)
    doc.add_paragraph()


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.85)
        s.bottom_margin = Inches(0.85)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    p(doc, "BITS Pilani WILP · PCAM* ZG516 Applied Generative AI", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "Group 15 — Demo / Presentation Speaking Script", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(
        doc,
        "BANKING77 Banking Customer-Service Assistant",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
    )
    p(
        doc,
        "Target length: 8–10 minutes  |  Notebook: BANKING77_GenAI_Assistant.ipynb  |  "
        "Repo: https://github.com/bits-pgaiml-agai/pcam-zg516-a1-banking77",
        size=10,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=12,
    )

    doc.add_heading("How to use this document", level=1)
    p(doc, "Each person reads only their section aloud. Paraphrase is fine; keep numbers exact. Grey italic lines are on-screen cues, not spoken.")
    p(doc, "If GPU/Colab fails: open evidence files and narrate the same words.")

    doc.add_heading("Speaking order (3 members)", level=1)
    table(
        doc,
        ["Time", "Speaker", "Segment"],
        [
            ["0:00–0:45", "Venkat (or Anil)", "Intro — problem, contracts, safety"],
            ["0:45–2:15", "Anil Kumar Reddy K", "Dataset, notebook, E1 baseline"],
            ["2:15–4:15", "Venkata Srinivasa Rao Tunuguntla", "E2 prompts + E3 decoding"],
            ["4:15–5:45", "Nisha Palanisamy", "E4 schema + E5 tool"],
            ["5:45–8:30", "Venkata Srinivasa Rao Tunuguntla", "Architecture, E6, metrics, E7"],
            ["8:30–10:00", "All (Venkat close)", "Limits, responsible use, Q&A"],
        ],
    )

    doc.add_heading("Shared facts (everyone should know)", level=1)
    p(doc, "Say if asked. Do not invent higher accuracy than committed evidence.")
    table(
        doc,
        ["Item", "Value"],
        [
            ["Dataset", "BANKING77 (PolyAI), seed 42"],
            ["Model", "Qwen/Qwen2.5-1.5B-Instruct (local E2–E5: Ollama qwen2.5:1.5b-instruct)"],
            ["JSON keys", "intent, response, action_required, tool_name, tool_arguments"],
            ["Tool", "transaction_status_lookup only; IDs TXN-DEMO-###"],
            ["Adaptation", "Prompting only (RAG/PEFT discussed, not implemented)"],
            ["Metrics (separate)", "Intent · Schema · Safety · Tool (+ latency)"],
        ],
    )
    p(doc, "Pipeline: Customer text → prompt → LLM → validate JSON → tool or safe fallback.", bold=True)

    # ----- INTRO -----
    doc.add_page_break()
    doc.add_heading("1. INTRO — Venkat or Anil (~45 seconds)", level=1)
    cue(doc, "[Camera / share screen: title slide or repo README]")
    speak(
        doc,
        "Hello. We are Group 15. Our project is a banking customer-service assistant on the "
        "BANKING77 dataset. The assistant must understand the customer intent among seventy-seven "
        "labels, write a safe reply, return strict JSON, and optionally call one mock tool for "
        "transaction status.",
    )
    speak(
        doc,
        "We did not train a large language model from scratch. We use a pretrained Qwen 2.5 "
        "Instruct model with prompting, validation, and a synthetic tool. This is an academic "
        "simulation — not a real bank — and we never handle real credentials.",
    )

    # ----- ANIL -----
    doc.add_heading("2. ANIL KUMAR REDDY K — Data, notebook, E1 (~1.5 minutes)", level=1)
    cue(doc, "[Open: data/eval/development_evaluation.csv, then evidence/baseline/E1_summary.json]")
    speak(
        doc,
        "Hi, I am Anil. I own data preparation, the Colab notebook backbone, and the E1 baseline.",
    )
    speak(
        doc,
        "We use the public PolyAI BANKING77 train and test CSVs, with a fixed seed of forty-two. "
        "We built frozen evaluation slices: a forty-row development set for design, and a reserved "
        "official test pilot that we do not use for prompt tuning.",
    )
    speak(
        doc,
        "The shared notebook is BANKING77_GenAI_Assistant. It loads the pinned Hugging Face model "
        "Qwen 2.5 1.5B Instruct, applies the chat template, and runs a simple baseline prompt with "
        "greedy decoding.",
    )
    speak(
        doc,
        "For E1 we ran three development examples. Result: zero of three outputs passed strict "
        "schema validation, and zero validated intents were correct. So E1 is an honest weak "
        "baseline, not a performance claim.",
    )
    speak(
        doc,
        "In the raw outputs we already see later failure modes: Markdown-wrapped JSON, invented "
        "intent labels, and unsafe requests for card details with a non-allowlisted tool. My "
        "deliverables are under evidence/baseline/, the data folders, and the notebook.",
    )
    cue(doc, "[Optional click: one bad line in E1_baseline_outputs.jsonl]")

    # ----- VENKAT E2/E3 -----
    doc.add_heading("3. VENKAT — E2 prompts and E3 decoding (~2 minutes)", level=1)
    cue(doc, "[Open: evidence/prompts_decoding/E2_summary.json and E3_summary.json]")
    speak(
        doc,
        "I am Venkat. I also own the E2 prompt-engineering and E3 decoding experiments.",
    )
    speak(
        doc,
        "For E2 we compared four prompt variants on the same ten development utterances, all "
        "with greedy decoding: zero-shot, few-shot, role-context, and constraint-based. Every "
        "variant includes the seventy-seven intent names, the JSON contract, and privacy rules.",
    )
    speak(
        doc,
        "On local Ollama qwen 2.5 1.5b-instruct, zero-shot was best with JSON valid rate zero "
        "point six zero, and validated intent correctness zero point one zero. Longer "
        "constraint and role prompts did not help this small local model. We lock zero-shot "
        "for this evidence pack. We do not claim seventy-seven-class accuracy — n equals ten.",
    )
    speak(
        doc,
        "For E3 we fixed zero-shot and compared greedy, temperature zero point seven with "
        "top-k fifty, and temperature zero point seven with top-p zero point nine. Prefer "
        "greedy for structured JSON. Evidence is under evidence/prompts_decoding/ and report "
        "sections seven and eight.",
    )

    # ----- NISHA -----
    doc.add_heading("4. NISHA PALANISAMY — E4 schema and E5 tool (~1.5 minutes)", level=1)
    cue(doc, "[Open: schemas/response.schema.json, then run dispatch_tool demos]")
    speak(
        doc,
        "Hi, I am Nisha. I own structured output validation, prompt templates, and the mock tool.",
    )
    speak(
        doc,
        "Our machine-readable contract has five keys: intent, response, action_required, "
        "tool_name, and tool_arguments. Intent must be an exact BANKING77 label. The only "
        "allowed tool is transaction_status_lookup, or null when no tool is needed.",
    )
    speak(
        doc,
        "The validator rejects Markdown fences, invented intents, bad tools, and bad "
        "transaction ID patterns. Invalid JSON never reaches the tool dispatcher. For E4, "
        "ten of ten designed accept and reject cases matched expectations.",
    )
    speak(
        doc,
        "For E5, the tool is pure Python with demo IDs like TXN-DEMO-001. Valid IDs return "
        "pending or completed with simulation true. Bad formats, unknown IDs, and wrong tool "
        "names are rejected. In the live LLM path, when the model invented a non-BANKING77 "
        "intent, validation blocked the tool and we returned a safe escalation message.",
    )
    cue(
        doc,
        '[Live] dispatch_tool("transaction_status_lookup", {"transaction_id": "TXN-DEMO-001"}) '
        "then try TXN-999 — show rejection.",
    )

    # ----- VENKAT ARCH/E6/E7 -----
    doc.add_heading("5. VENKAT — Architecture, E6, metrics, E7 (~2.5–3 minutes)", level=1)
    cue(doc, "[Open: architecture diagram / report §5, then E6_failure_analysis.csv]")
    speak(
        doc,
        "Continuing — architecture and evaluation. We compared causal LLMs, VAEs, and GANs, "
        "with diffusion as a contrast. We selected a causal LLM because we need intent, natural "
        "language, constrained JSON, and optional tool requests in one prompted workflow.",
    )
    speak(
        doc,
        "Architecture: customer utterance goes to a prompt builder, then the model, then "
        "strict JSON validation. Invalid outputs refuse tools and escalate. Valid outputs "
        "return the response or call only transaction_status_lookup on TXN-DEMO IDs.",
    )
    speak(
        doc,
        "Adaptation: we compared prompting, RAG, and PEFT. For this prototype we implemented "
        "prompting plus validation plus an allowlisted tool. RAG and PEFT remain future options.",
    )
    speak(
        doc,
        "We report four metrics separately: intent correctness, schema validity, safety, and "
        "tool-call correctness — never one unsupported accuracy number. E1 was zero of three "
        "schema-valid. Best local E2 schema rate was zero point six zero on ten rows.",
    )
    speak(
        doc,
        "E6 failures: F1 Markdown-fenced JSON; F2 invented intent label; F3 wrong intent plus "
        "asking for a card number with a non-allowlisted tool. Mitigations are constraint "
        "prompts, exact labels, validators, privacy rules, and tool allowlisting.",
    )
    speak(
        doc,
        "E7 is the end-to-end path: build prompt, generate, validate, tool or fallback, log "
        "metrics. In this demo we show a failing baseline case, a valid JSON path or evidence, "
        "a successful mock tool call, and we state the four metrics separately.",
    )

    # ----- CLOSE -----
    doc.add_heading("6. CLOSE — All / Venkat (~1–1.5 minutes)", level=1)
    speak(
        doc,
        "To close: this is a simulation only. We do not connect to a real bank, and we never "
        "solicit PINs, OTPs, or card numbers. Sample sizes are small, so we do not claim full "
        "BANKING77 benchmark accuracy. Future work includes Colab re-run on the pinned Hugging "
        "Face checkpoint and expanding the development pilot. Thank you — we are happy to take "
        "questions.",
    )

    doc.add_heading("Quick Q&A cheat sheet", level=1)
    table(
        doc,
        ["If asked…", "Answer…"],
        [
            ["Did you fine-tune?", "No — prompting only."],
            ["Why causal LLM?", "Generation + JSON + tools; classifier alone is incomplete."],
            ["Accuracy claim?", "No single accuracy; report rates with exact n and split."],
            ["Real bank data?", "No — public BANKING77 + mock TXN-DEMO only."],
            ["Best prompt (local)?", "zero_shot on qwen2.5:1.5b-instruct (JSON 0.60, n=10)."],
            ["Decoding?", "Prefer greedy for structured JSON."],
            ["Who did E2/E3?", "Venkat (reassigned after a member left; team of three)."],
        ],
    )

    doc.add_heading("On-screen checklist during recording", level=1)
    for item in [
        "☐ Show failing E1 raw output / validator rejection",
        "☐ Show E2 comparison numbers (or E2_summary.json)",
        "☐ Show E3 decoding table",
        "☐ Show schema + tool success TXN-DEMO-001",
        "☐ Show tool rejection on bad ID",
        "☐ State four metrics separately",
        "☐ Say: simulation / no real bank credentials",
    ]:
        p(doc, item, size=11, after=4)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
