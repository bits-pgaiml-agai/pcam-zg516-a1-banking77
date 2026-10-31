"""Build the merged Group 15 BANKING77 technical report (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "BANKING77_Group15_Technical_Report.docx"
BASELINE_DOCX = ROOT / "reports" / "BANKING77_Report_Sections_2_3_6_Problem_Dataset_Baseline.docx"


def set_run_font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, *, bold=False, italic=False, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def copy_table(src_table, doc):
    rows = [[c.text.strip() for c in row.cells] for row in src_table.rows]
    if not rows:
        return
    add_table(doc, rows[0], rows[1:])


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # ----- Title page -----
    for _ in range(2):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("BITS Pilani — Work Integrated Learning Programmes")
    set_run_font(r, size=14, bold=True)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PCAM* ZG516 Applied Generative AI")
    set_run_font(r, size=13, bold=True)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Group Major Project / Assignment-1")
    set_run_font(r, size=12)

    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "BANKING77 Banking Customer-Service Assistant:\n"
        "Prompted Causal LLM with Structured Output and Mock Tool Calling"
    )
    set_run_font(r, size=16, bold=True)

    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Dataset D1: BANKING77 (PolyAI)")
    set_run_font(r, size=12)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Group 15")
    set_run_font(r, size=14, bold=True)

    doc.add_paragraph()
    add_table(
        doc,
        ["Name", "Workstream / contribution focus"],
        [
            ["Anil Kumar Reddy K", "Data prep, E1 baseline, notebook backbone (report §§2, 3, 6)"],
            ["Nisha Palanisamy", "E4 structured output, E5 tools, templates (report §10)"],
            [
                "Venkata Srinivasa Rao Tunuguntla",
                "E2/E3 prompts & decoding (§§7–8); architecture, adaptation, E6/E7, report merge (§§4, 5, 9, 11–14)",
            ],
        ],
    )
    add_para(doc, "Repository: https://github.com/bits-pgaiml-agai/pcam-zg516-a1-banking77", italic=True)
    add_para(
        doc,
        "Primary local E2–E5 runtime: Ollama qwen2.5:1.5b-instruct "
        "(aligned to HF pin Qwen/Qwen2.5-1.5B-Instruct). Seed: 42.",
        italic=True,
    )
    doc.add_page_break()

    # ----- 1 Abstract -----
    add_heading(doc, "1. Abstract", 1)
    add_para(
        doc,
        "This project builds a banking customer-service assistant on the public BANKING77 "
        "intent dataset. A pretrained causal instruct LLM (Qwen2.5-1.5B-Instruct / local "
        "Ollama qwen2.5:1.5b-instruct) is adapted by prompting only—no training from scratch "
        "and no weight updates. The system maps a customer utterance to a strict JSON object "
        "with intent, a safe natural-language response, and optional allowlisted tool fields, "
        "then may call a synthetic transaction_status_lookup backend using demo IDs only.",
    )
    add_para(
        doc,
        "Required experiments E1–E7 are completed: a weak E1 baseline (0/3 schema-valid), "
        "four prompt variants (best local JSON validity 0.60 with zero_shot), decoding "
        "comparison (greedy preferred for structured JSON), schema validation (10/10 designed "
        "cases), mock tool invocation with blocked unsafe/invalid calls, failure analysis "
        "(≥3 observed cases), and an integrated notebook/demo workflow. Metrics are reported "
        "separately (intent, schema, safety, tool, latency) with explicit sample sizes; we do "
        "not claim 77-class accuracy. The application is an academic simulation and must not "
        "handle real banking credentials or personal financial data.",
    )

    # ----- Pull §§2,3,6 from baseline docx -----
    base = Document(str(BASELINE_DOCX))
    base_tables = list(base.tables)
    bp = [p.text.strip() for p in base.paragraphs if p.text.strip()]

    add_heading(doc, "2. Problem statement and motivation", 1)
    add_heading(doc, "2.1 User problem and application scope", 2)
    add_para(doc, bp[5])
    add_para(doc, bp[6])
    add_heading(doc, "2.2 Why a generative model is appropriate", 2)
    add_para(doc, bp[8])
    add_heading(doc, "2.3 Success criteria and limits", 2)
    add_para(doc, bp[10])
    add_para(doc, bp[11])

    add_heading(doc, "3. Dataset description and preparation", 1)
    add_heading(doc, "3.1 Source and data fields", 2)
    add_para(doc, bp[14])
    add_para(doc, bp[15])
    copy_table(base_tables[0], doc)
    add_heading(doc, "3.2 Observed dataset properties", 2)
    copy_table(base_tables[1], doc)
    add_heading(doc, "3.3 Representative development utterances", 2)
    add_para(doc, bp[18])
    add_para(doc, bp[19])
    add_para(doc, bp[20])
    add_para(doc, bp[21])
    add_heading(doc, "3.4 Preprocessing procedure", 2)
    add_para(doc, bp[24])
    add_para(doc, bp[25])
    add_heading(doc, "3.5 Leakage controls and coverage", 2)
    add_para(doc, bp[27])
    add_para(doc, bp[28])
    copy_table(base_tables[2], doc)
    add_heading(doc, "3.6 Reproducibility records", 2)
    add_para(doc, bp[30])

    # ----- 4 Architecture analysis -----
    add_heading(doc, "4. Generative-model and LLM architecture analysis", 1)
    add_heading(doc, "4.1 Generative-model family comparison", 2)
    add_para(
        doc,
        "At least three generative-model families from the course are relevant. We select a "
        "pretrained causal / autoregressive LLM for this application.",
    )
    add_table(
        doc,
        ["Family", "How it generates", "Fit for BANKING77 assistant"],
        [
            [
                "Causal / autoregressive LLM (selected)",
                "Next-token prediction conditioned on prompt; chat templates support instructions",
                "Direct fit for intent + natural-language reply + optional tool request in one prompted workflow",
            ],
            [
                "VAE",
                "Encode to a latent distribution; sample/decode",
                "Strong for representation learning; weak fit for constrained multi-field JSON and policy-safe wording without major redesign",
            ],
            [
                "GAN",
                "Generator vs discriminator adversarial training",
                "Useful for synthetic data realism; not a natural controller for allowlisted tools, refusals, or exact intent labels",
            ],
            [
                "Diffusion (contrast only)",
                "Iterative denoising from noise",
                "Dominant for continuous media; not required for this short text JSON assistant",
            ],
        ],
    )
    add_para(
        doc,
        "Decision: use a pretrained causal LLM. The need is language understanding plus "
        "constrained generation, not image synthesis or latent reconstruction. VAE/GAN/diffusion "
        "remain comparison points; this project does not train them.",
    )
    add_heading(doc, "4.2 LLM building blocks in this project", 2)
    add_table(
        doc,
        ["Component", "Role in the banking assistant"],
        [
            ["Tokenizer", "Turns system rules, 77 intent names, optional few-shot examples, and the utterance into token IDs via the chat template"],
            ["Embeddings", "Map token IDs into vectors for the Transformer"],
            ["Transformer / attention", "Conditions generation on the full prompt (labels, safety rules, request)"],
            ["Language-modelling head", "Scores the next-token distribution"],
            ["Autoregressive decode", "Prefill processes the prompt; decode emits JSON fields until EOS or max_new_tokens"],
            ["Decoding controls", "Greedy for stable baselines; temperature / top-k / top-p in E3 to study variability vs validity"],
        ],
    )
    add_para(
        doc,
        "Implemented model: Qwen/Qwen2.5-1.5B-Instruct (pinned revision), Hugging Face "
        "AutoModelForCausalLM, with local Ollama qwen2.5:1.5b-instruct used for E2–E5 evidence "
        "when HF weights were unavailable. No training from scratch and no weight updates.",
    )
    add_heading(doc, "4.3 Why not just a classifier?", 2)
    add_para(
        doc,
        "A classical intent classifier could score the 77 labels well. The assignment also "
        "requires response generation, a decoding study, structured output, and a tool/function "
        "workflow. A prompted causal LLM covers that end-to-end path; classification accuracy "
        "alone would underspecify the deliverable.",
    )

    # ----- 5 Application architecture -----
    add_heading(doc, "5. Proposed application architecture", 1)
    add_para(
        doc,
        "End-to-end control flow:",
        bold=True,
    )
    add_para(
        doc,
        "Customer utterance → Prompt builder (baseline / zero-shot / few-shot / role / "
        "constraint) → Tokenizer + chat template → Qwen2.5-1.5B-Instruct → decoding → "
        "Strict JSON extract + validate. Invalid outputs refuse tools and return a safe "
        "fallback. Valid outputs either return the response or, if action_required, call "
        "allowlisted transaction_status_lookup on TXN-DEMO-### IDs, then produce a grounded "
        "demo reply or escalation text.",
    )
    add_bullets(
        doc,
        [
            "Output keys: intent, response, action_required, tool_name, tool_arguments",
            "Tool: transaction_status_lookup only; synthetic IDs TXN-DEMO-###",
            "Development pilot (40) for prompt/decoding design; official test pilot (40) reserved after lock; seed 42",
            "Modules: notebooks/BANKING77_GenAI_Assistant.ipynb, schemas/, src/tools.py, data/eval/",
        ],
    )

    # ----- 6 Baseline -----
    add_heading(doc, "6. Baseline LLM implementation", 1)
    add_heading(doc, "6.1 Model and inference workflow", 2)
    add_para(doc, bp[33])
    add_para(doc, bp[34])
    add_heading(doc, "6.2 Output validation", 2)
    add_para(doc, bp[36])
    add_heading(doc, "6.3 Three observed baseline generations", 2)
    copy_table(base_tables[3], doc)
    add_para(doc, bp[38])
    add_para(doc, bp[39])
    add_para(doc, bp[40])
    add_para(doc, bp[41])
    add_heading(doc, "6.4 Evidence and further evaluation", 2)
    add_para(doc, bp[43])
    add_para(
        doc,
        "E1 summary (evidence/baseline/E1_summary.json): n=3; schema_valid=0; "
        "intent_correct=0; general_accuracy_claim=false.",
    )
    copy_table(base_tables[4], doc)

    # ----- 7 Prompt engineering -----
    add_heading(doc, "7. Prompt-engineering experiments", 1)
    add_para(
        doc,
        "Primary group HF contract: Qwen/Qwen2.5-1.5B-Instruct. Local E2/E3 evidence used "
        "Ollama qwen2.5:1.5b-instruct. An exploratory comparison with qwen2.5:3b is in the "
        "appendix; primary figures below are 1.5b-instruct only.",
    )
    add_heading(doc, "7.1 Design", 2)
    add_table(
        doc,
        ["Item", "Choice"],
        [
            ["Variants (4)", "zero_shot, few_shot, role_context, constraint_based"],
            ["Shared rules", "Exact 77-label list + JSON contract + privacy rules"],
            ["Few-shot", "4 fixed examples from data/eval/in_prompt_examples.csv"],
            ["Evaluation subset", "10 rows from development_evaluation.csv, seed 42"],
            ["Decoding for E2", "Greedy / temperature 0"],
            ["Runtime", "Ollama qwen2.5:1.5b-instruct"],
        ],
    )
    add_heading(doc, "7.2 Results (n=10 development subset)", 2)
    add_table(
        doc,
        ["Variant", "JSON valid rate", "Intent correct (validated)", "Intent accuracy", "Mean latency (s)"],
        [
            ["zero_shot (best)", "0.60", "0.10", "0.10", "8.3"],
            ["few_shot", "0.30", "0.10", "0.10", "6.6"],
            ["role_context", "0.30", "0.10", "0.10", "8.4"],
            ["constraint_based", "0.10", "0.00", "0.00", "9.7"],
        ],
    )
    add_bullets(
        doc,
        [
            "All improved prompts beat E1’s 0% schema acceptance on at least some rows; zero_shot was strongest (0.60 JSON valid).",
            "Validated intent correctness remained low (0.10 best) on this 1.5B local run.",
            "Longer constraint/role prompts reduced schema validity vs zero-shot on this checkpoint.",
            "n=10 covers 10 intents only; do not claim 77-class performance.",
            "Recommended lock for this evidence pack: zero_shot.",
        ],
    )

    # ----- 8 Decoding -----
    add_heading(doc, "8. Decoding experiments", 1)
    add_heading(doc, "8.1 Design", 2)
    add_table(
        doc,
        ["Item", "Choice"],
        [
            ["Prompt", "zero_shot (best E2 variant on this runtime)"],
            ["Settings", "greedy (T=0); T=0.7 + top-k=50; T=0.7 + top-p=0.9"],
            ["Rows", "3 utterances from the E2 subset (seed 43)"],
        ],
    )
    add_heading(doc, "8.2 Results (n=3 × 3 settings)", 2)
    add_table(
        doc,
        ["Decoding", "JSON valid rate", "Intent correct rate", "Mean latency (s)"],
        [
            ["greedy (T=0)", "0.33", "0.00", "6.3"],
            ["T=0.7 + top-k=50", "0.00", "0.00", "5.9"],
            ["T=0.7 + top-p=0.9", "0.33", "0.00", "5.8"],
        ],
    )
    add_para(
        doc,
        "For production-like JSON, keep greedy after locking the prompt; retain sampling "
        "rows as the required decoding study. Intent correctness was 0 on this tiny sample "
        "(observational only).",
    )

    # ----- 9 Adaptation -----
    add_heading(doc, "9. Adaptation strategy: prompting vs RAG vs PEFT", 1)
    add_table(
        doc,
        ["Approach", "What it changes", "Fit for this 2-week prototype", "Cost / risk"],
        [
            [
                "Prompting (selected)",
                "Instructions, examples, constraints, decoding",
                "Matches E1–E5; reproducible on Colab; no training",
                "Prompt brittleness; label invention; JSON wrappers (seen in E1)",
            ],
            [
                "RAG",
                "Retrieve approved FAQ/policy snippets into context",
                "Useful if answers must cite current bank policy",
                "Needs curated corpus; retrieval errors can look authoritative",
            ],
            [
                "PEFT / LoRA",
                "Small adapters on frozen base weights",
                "Could improve JSON/intent consistency with labelled pairs",
                "Extra compute, data prep, eval burden; optional unless approved",
            ],
        ],
    )
    add_para(
        doc,
        "Implemented choice: prompting (+ strict validation + allowlisted tool). RAG and "
        "PEFT are documented trade-offs for future work, not required for the baseline submission path.",
    )

    # ----- 10 Structured + tools -----
    add_heading(doc, "10. Structured output and tool/function calling", 1)
    add_heading(doc, "10.1 Output schema", 2)
    add_table(
        doc,
        ["Key", "Type", "Rule"],
        [
            ["intent", "string", "Exact BANKING77 label from the 77-name list"],
            ["response", "string", "Non-empty customer-facing text"],
            ["action_required", "boolean", "Whether a tool should run"],
            ["tool_name", "string | null", "Only transaction_status_lookup or null"],
            ["tool_arguments", "object", "{} if no tool; else {transaction_id: TXN-DEMO-###}"],
        ],
    )
    add_para(
        doc,
        "Strict parser rejects Markdown fences, prose around JSON, unknown keys, invented "
        "intents, non-boolean action_required, and non-allowlisted tools. Invalid outputs "
        "never reach dispatch_tool.",
    )
    add_heading(doc, "10.2 E4 validation evidence", 2)
    add_para(doc, "Designed cases: 10/10 matched expectations (evidence/structured_tools/E4_summary.json).")
    add_table(
        doc,
        ["Case", "Expect", "Result"],
        [
            ["V1 valid no tool", "accept", "accept"],
            ["V2 valid tool request", "accept", "accept"],
            ["R1 Markdown fence", "reject", "reject"],
            ["R2 Invented intent", "reject", "reject"],
            ["R3 action_required as string", "reject", "reject"],
            ["R4 Unallowlisted tool", "reject", "reject"],
            ["R5 action_required=true with null tool", "reject", "reject"],
            ["R6 Prose before JSON", "reject", "reject"],
            ["R7 Extra keys", "reject", "reject"],
            ["R8 Bad transaction id pattern", "reject", "reject"],
        ],
    )
    add_heading(doc, "10.3 Programmatic prompting", 2)
    add_para(
        doc,
        "Reusable Jinja2 template prompts/assistant_template.jinja2 and Python helper "
        "src/prompt_template.py render messages from customer_text, intent_names, variant, "
        "include_tool, and optional few_shot_block.",
    )
    add_heading(doc, "10.4 Tool / function calling (E5)", 2)
    add_para(
        doc,
        "Tool transaction_status_lookup is implemented in src/tools.py as a synthetic "
        "dictionary only (no bank network). Schema: schemas/tool_transaction_status_lookup.json.",
    )
    add_table(
        doc,
        ["Case", "Outcome"],
        [
            ["TXN-DEMO-001", "ok, status=pending, simulation=true"],
            ["TXN-DEMO-002", "ok, status=completed"],
            ["TXN-DEMO-999", "not found"],
            ["TXN-1001", "invalid format"],
            ["verify_my_identity", "not allowlisted"],
            ["Extra arguments", "rejected"],
        ],
    )
    add_para(
        doc,
        "Integrated live LLM attempt on qwen2.5:1.5b-instruct (zero_shot): invented "
        "non-BANKING77 intent and inconsistent action/tool fields → validator blocked tool "
        "(safe failure). Representative valid pipeline demonstrates schema-valid tool request → "
        "mock lookup → grounded demo reply with simulation disclosed.",
    )

    # ----- 11 Evaluation -----
    add_heading(doc, "11. Evaluation and results", 1)
    add_heading(doc, "11.1 Criteria", 2)
    add_para(doc, "Criteria are reported separately and never collapsed into one unsupported accuracy number.")
    add_table(
        doc,
        ["#", "Criterion", "How measured"],
        [
            ["1", "Intent correctness", "Exact match of validated intent vs gold; invalid JSON counts as incorrect for the pipeline"],
            ["2", "Schema validity", "Strict parser + contract checks"],
            ["3", "Response safety / relevance", "Manual rubric 1–5 + binary no-credential / no-fake-completion"],
            ["4", "Tool-call correctness", "Allowlist, argument shape, ID pattern; no call on invalid JSON"],
            ["5", "Latency / tokens", "Wall time and token counts (ops)"],
        ],
    )
    add_heading(doc, "11.2 Comparative results", 2)
    add_table(
        doc,
        ["System", "Schema valid rate", "Intent accuracy (validated)", "Notes"],
        [
            ["E1 baseline", "0/3", "0/3", "Markdown / invented label / credential ask"],
            ["Best prompt E2 zero_shot", "0.60 (n=10)", "0.10", "Beats E1 schema; intent still weak on 1.5B local"],
            ["Best decoding E3 greedy", "0.33 (n=3)", "0.00", "Prefer greedy; top-k produced 0 valid JSON"],
            ["E4/E5 validation+tool", "E4 designed 10/10", "N/A (controls)", "Invalid JSON blocked tool; mock TXN-DEMO OK"],
        ],
    )
    add_heading(doc, "11.3 Integrated final workflow (E7)", 2)
    add_para(
        doc,
        "The integrated demo follows evidence/evaluation/E7_integration_checklist.md and "
        "demo/VIDEO_LINK.txt. In the notebook we load pinned BANKING77 subsets, build "
        "messages with the locked prompt variant (zero_shot for the local 1.5b-instruct "
        "evidence pack), generate with greedy decoding, run validate_output, optionally "
        "dispatch_tool for TXN-DEMO-###, and log metrics against the rubric. The 8–10 minute "
        "group demonstration shows a failing E1 case, a valid JSON path, a successful mock "
        "tool call, a credential-bait refusal, and states the four primary metrics separately.",
    )
    add_para(
        doc,
        "Flags: RUN_TEAM_EXPERIMENTS controls E2–E5 generation cells; RUN_FINAL_EVAL must "
        "remain false until the prompt is locked on development data. If live GPU is "
        "unavailable, narrate from committed evidence under evidence/.",
    )

    # ----- 12 Failure analysis -----
    add_heading(doc, "12. Failure analysis and responsible-use considerations", 1)
    add_heading(doc, "12.1 E6 failure cases", 2)
    add_table(
        doc,
        ["ID", "Input (short)", "Gold", "Failure category", "Mitigation"],
        [
            [
                "F1",
                "Top-up reverted by app",
                "top_up_reverted",
                "Malformed structured output (Markdown fences; action without tool)",
                "Bare JSON only; reject fences; require tool iff action",
            ],
            [
                "F2",
                "Why can't I use disposable card?",
                "disposable_card_limits",
                "Label invention",
                "Few-shot + exact label copy; validator drops unknown intents",
            ],
            [
                "F3",
                "ATM didn't return card",
                "card_swallowed",
                "Unsafe + wrong intent + bad tool",
                "Privacy rules; allowlist only transaction_status_lookup; escalate",
            ],
            [
                "F4–F6",
                "Ambiguous transfer / missing TXN / PIN bait",
                "designed controls",
                "Ambiguity, missing args, credential bait",
                "Clarify/escalate; pattern-check IDs; refuse PIN collection",
            ],
        ],
    )
    add_heading(doc, "12.2 Responsible use", 2)
    add_bullets(
        doc,
        [
            "Academic simulation only — no bank connectivity, no real balances, blocks, or transfers.",
            "Never solicit or echo PINs, passwords, OTPs, or full card numbers.",
            "Do not invent fees, eligibility, or transaction status; label tool results as demo.",
            "Ambiguous or high-impact requests → escalate to the bank’s official channel.",
            "Valid JSON ≠ truthful or safe reply; keep manual review in the scoring pack.",
            "Retain BANKING77 CC BY 4.0 attribution (docs/baseline/ATTRIBUTION.txt).",
        ],
    )

    # ----- 13 Limitations -----
    add_heading(doc, "13. Limitations", 1)
    add_bullets(
        doc,
        [
            "40-row pilots cover 40/77 intents; E1 uses only 3 rows — not a full-benchmark claim.",
            "BANKING77 has no gold responses or tool labels — response/tool metrics need human review.",
            "Small instruct model can invent labels and wrap JSON; validation is mandatory.",
            "Prompting-only adaptation may plateau; RAG/PEFT not implemented in this iteration.",
            "Hardware and runtime (CPU/GPU/Ollama vs HF) can change generations; compare manifests.",
            "Local 1.5B instruct improves schema vs E1 but intent remains low.",
            "Live LLM→tool success still depends on producing schema-valid BANKING77 intents.",
        ],
    )

    # ----- 14 Conclusion -----
    add_heading(doc, "14. Conclusion and future work", 1)
    add_para(
        doc,
        "We delivered a prompting-based BANKING77 assistant with reproducible data prep, a "
        "documented weak baseline, prompt and decoding studies, strict structured-output "
        "validation, a mock tool workflow, failure analysis, and an integrated demo path. "
        "Primary local E2–E5 numbers use qwen2.5:1.5b-instruct to stay aligned with the "
        "official HF 1.5B Instruct pin.",
    )
    add_para(
        doc,
        "Future work: re-run E2–E5 on Colab with the pinned HF checkpoint; expand evaluation "
        "to the full 40-row development pilot (and reserved test only after lock); optionally "
        "add instructor-approved FAQ RAG or PEFT if prompting plateaus; strengthen safety "
        "review coverage for generated replies.",
    )

    # ----- Contribution statements -----
    add_heading(doc, "Individual contribution statements", 1)
    add_para(
        doc,
        "Anil Kumar Reddy K: responsible for BANKING77 dataset analysis/preparation, E1 "
        "baseline, notebook backbone, and report §§2, 3, 6.",
    )
    add_para(doc, "Member confirmation/signature: ____________________    Date: _______________")
    add_para(
        doc,
        "Nisha Palanisamy: responsible for E4 structured JSON schema/validation, programmatic "
        "prompt templates, E5 tool schema/invocation, and report §10.",
    )
    add_para(doc, "Member confirmation/signature: ____________________    Date: _______________")
    add_para(
        doc,
        "Venkata Srinivasa Rao Tunuguntla: responsible for E2 prompt-engineering experiments "
        "and E3 decoding study (report §§7–8); generative-model/LLM architecture analysis; "
        "application architecture; prompting vs RAG vs PEFT; evaluation rubric; E6 failure "
        "analysis; responsible-use/limitations; E7 integration/demo; conclusion; and merging "
        "the group report (§§4, 5, 9, 11–14).",
    )
    add_para(doc, "Member confirmation/signature: ____________________    Date: _______________")

    # ----- References -----
    add_heading(doc, "References", 1)
    refs = [
        "[1] PolyAI. BANKING77 dataset card and canonical label mapping. https://huggingface.co/datasets/PolyAI/banking77",
        "[2] Casanueva et al. Efficient Intent Detection with Dual Sentence Encoders, 2020. https://arxiv.org/abs/2003.04807",
        "[3] Qwen team. Qwen2.5-1.5B-Instruct model card. https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct",
        "[4] Hugging Face. Transformers documentation — chat templates. https://huggingface.co/docs/transformers/chat_templating",
        "[5] BITS Pilani WILP. PCAM* ZG516 Group Major Project Assignment 1 (CS1–CS5). Supplied assignment document.",
        "[6] Group 15 repository. https://github.com/bits-pgaiml-agai/pcam-zg516-a1-banking77",
    ]
    for ref in refs:
        add_para(doc, ref)

    # ----- Appendix -----
    doc.add_page_break()
    add_heading(doc, "Appendix", 1)
    add_heading(doc, "A. Prompt variant summary", 2)
    add_table(
        doc,
        ["Variant", "Instruction focus"],
        [
            ["baseline (E1 only)", "Simple classify + safe response; return JSON"],
            ["zero_shot", "Classify + safe response; no examples"],
            ["few_shot", "Same + 4 fixed in-prompt examples"],
            ["role_context", "First-line support persona; escalate account-specific actions"],
            ["constraint_based", "Emphasize bare JSON, exact keys/types, exact label copy"],
        ],
    )
    add_para(
        doc,
        "Exact system/user messages sent to the LLM are archived in "
        "prompts/ACTUAL_PROMPTS_E2.md and prompts/sample_rendered_*.json "
        "(full 77-label list included in each rendered sample).",
    )
    add_heading(doc, "B. Response schema (excerpt)", 2)
    add_para(
        doc,
        'Required keys: intent, response, action_required, tool_name, tool_arguments. '
        "additionalProperties=false. Full file: schemas/response.schema.json.",
        italic=True,
    )
    add_heading(doc, "C. Sample validated / rejected outputs", 2)
    add_para(
        doc,
        "Accept examples: V1 valid no-tool JSON; V2 valid tool request with "
        'transaction_id "TXN-DEMO-001". Reject examples: Markdown-fenced JSON (R1), invented '
        "intent (R2), unallowlisted tool (R4), bad TXN pattern (R8). See "
        "evidence/structured_tools/E4_validation_cases.csv.",
    )
    add_heading(doc, "D. Tool specification (summary)", 2)
    add_para(
        doc,
        "Name: transaction_status_lookup. Argument: transaction_id matching ^TXN-DEMO-[0-9]{3}$. "
        "Returns status from a fixed mock table or a simulation error. Implementation: src/tools.py.",
    )
    add_heading(doc, "E. Exploratory Ollama 1.5b-instruct vs 3b comparison (not primary)", 2)
    add_para(
        doc,
        "Primary submission figures use qwen2.5:1.5b-instruct. The table below records an "
        "exploratory side-by-side under the same protocol (informational only).",
    )
    add_table(
        doc,
        ["Variant", "1.5b-instruct JSON/intent", "3b JSON/intent"],
        [
            ["zero_shot", "0.60 / 0.10", "0.50 / 0.20"],
            ["few_shot", "0.30 / 0.10", "0.60 / 0.30"],
            ["role_context", "0.30 / 0.10", "0.70 / 0.20"],
            ["constraint_based", "0.10 / 0.00", "0.60 / 0.30"],
        ],
    )
    add_para(
        doc,
        "Takeaway: 3b scored higher, but the group standardizes on 1.5b-instruct to match the "
        "HF 1.5B Instruct pin. Full note: docs/MODEL_COMPARISON_OLLAMA.md.",
    )
    add_heading(doc, "F. Code and evidence index", 2)
    add_bullets(
        doc,
        [
            "Notebook: notebooks/BANKING77_GenAI_Assistant.ipynb",
            "E2/E3 runners: scripts/run_e2_e3.py, scripts/run_e2_e3_ollama.py",
            "E4/E5 runner: scripts/run_e4_e5.py",
            "Evidence: evidence/baseline/, evidence/prompts_decoding/, evidence/structured_tools/, evidence/evaluation/",
            "Demo video: demo/VIDEO_LINK.txt",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
